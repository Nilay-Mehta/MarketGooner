import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

# Brand Palette (Modern Executive Tech / Venture Theme)
NAVY = RGBColor(26, 54, 93)       # #1A365D - Deep Slate Navy
CHARCOAL = RGBColor(45, 55, 72)   # #2D3748 - Dark Body Gray
MUTED = RGBColor(113, 128, 150)   # #718096 - Secondary Gray
ACCENT_BLUE = RGBColor(49, 130, 206) # #3182CE - Vivid Blue Accent
BORDER_GRAY = "E2E8F0"
BOX_BG_GRAY = "F7FAFC"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="KEY TAKEAWAY"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    # Left border only (accent blue), background light gray
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{BOX_BG_GRAY}"/>')
    cell._tc.get_or_add_tcPr().append(shading)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="3182CE"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title.upper()}\n")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(9.5)
    run_title.font.bold = True
    run_title.font.color.rgb = ACCENT_BLUE
    
    run_body = p.add_run(text)
    run_body.font.name = "Segoe UI"
    run_body.font.size = Pt(10)
    run_body.font.color.rgb = CHARCOAL
    run_body.font.italic = True
    
    # Space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

def format_inline_runs(paragraph, text, base_font="Segoe UI", base_size=10.5, base_color=CHARCOAL):
    """Parses basic inline markdown (**bold**, *italic*, `code`)."""
    # Regex splits while keeping matches
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.bold = True
            run.font.color.rgb = NAVY
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.italic = True
            run.font.color.rgb = base_color
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(f" {token[1:-1]} ")
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = ACCENT_BLUE
        else:
            run = paragraph.add_run(token)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color

def convert_md_to_docx(md_path: str, docx_path: str = None) -> str:
    md_file = Path(md_path)
    if not md_file.exists():
        raise FileNotFoundError(f"Markdown file not found: {md_path}")
        
    if not docx_path:
        docx_path = md_file.with_suffix('.docx')
        
    doc = Document()
    
    # Configure 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    # Parse YAML frontmatter if present
    frontmatter = {}
    if content.startswith('---'):
        parts = content.split('---', 2)
        if len(parts) >= 3:
            fm_text = parts[1]
            content = parts[2]
            for line in fm_text.strip().split('\n'):
                if ':' in line and not line.strip().startswith('-'):
                    k, v = line.split(':', 1)
                    frontmatter[k.strip()] = v.strip().strip('"').strip("'")

    # If frontmatter exists, generate a sleek metadata header table
    if frontmatter:
        meta_table = doc.add_table(rows=len(frontmatter), cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        
        for idx, (k, v) in enumerate(frontmatter.items()):
            row = meta_table.rows[idx]
            cell_k, cell_v = row.cells[0], row.cells[1]
            cell_k.width = Inches(2.0)
            cell_v.width = Inches(4.5)
            
            p_k = cell_k.paragraphs[0]
            p_k.paragraph_format.space_before = Pt(2)
            p_k.paragraph_format.space_after = Pt(2)
            run_k = p_k.add_run(k.replace('_', ' ').title())
            run_k.font.name = "Segoe UI"
            run_k.font.size = Pt(9.5)
            run_k.font.bold = True
            run_k.font.color.rgb = MUTED
            
            p_v = cell_v.paragraphs[0]
            p_v.paragraph_format.space_before = Pt(2)
            p_v.paragraph_format.space_after = Pt(2)
            run_v = p_v.add_run(str(v))
            run_v.font.name = "Segoe UI"
            run_v.font.size = Pt(9.5)
            run_v.font.color.rgb = CHARCOAL
            
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(8)
        p_div.paragraph_format.space_after = Pt(12)

    lines = content.split('\n')
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        # Parse table rows
        rows_data = []
        for tl in table_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*$', tl):
                continue # Skip separator row
            cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if cells:
                rows_data.append(cells)
                
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True
            
            for r_idx, r_data in enumerate(rows_data):
                row = tbl.rows[r_idx]
                is_header = (r_idx == 0)
                for c_idx in range(num_cols):
                    cell = row.cells[c_idx]
                    val = r_data[c_idx] if c_idx < len(r_data) else ""
                    
                    # Style cell
                    if is_header:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EDF2F7"/>')
                        cell._tc.get_or_add_tcPr().append(shd)
                        
                    borders = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_GRAY}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E0"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>'
                    )
                    cell._tc.get_or_add_tcPr().append(borders)
                    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    format_inline_runs(p, val, base_font="Segoe UI", base_size=9.5, base_color=NAVY if is_header else CHARCOAL)
                    if is_header:
                        for r in p.runs:
                            r.font.bold = True
            
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(6)
            p_space.paragraph_format.space_after = Pt(6)
            
        in_table = False
        table_lines = []

    for line in lines:
        line_str = line.strip()
        
        # Check for Markdown tables
        if line_str.startswith('|') and line_str.endswith('|'):
            in_table = True
            table_lines.append(line_str)
            continue
        elif in_table:
            flush_table()
            
        if not line_str:
            continue
            
        # H1 Title
        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line_str[2:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = NAVY
            
        # H2 Section Header
        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line_str[3:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = NAVY
            
        # H3 Subsection Header
        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line_str[4:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = ACCENT_BLUE
            
        # Blockquote / Highlight Callout
        elif line_str.startswith('> '):
            quote_text = line_str[2:].strip()
            add_callout_box(doc, quote_text, title="OPPORTUNITY INSIGHT")
            
        # Bullet Points
        elif line_str.startswith('- ') or line_str.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, line_str[2:].strip())
            
        # Numbered List
        elif re.match(r'^\d+\.\s+', line_str):
            num_match = re.match(r'^\d+\.\s+', line_str)
            text_body = line_str[num_match.end():].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, text_body)
            
        # Horizontal Rule
        elif line_str in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 50)
            run.font.color.rgb = RGBColor(226, 232, 240)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        # Regular Body Paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, line_str)
            
    if in_table:
        flush_table()
        
    doc.save(str(docx_path))
    print(f"✅ Successfully converted: {md_path} -> {docx_path}")
    return str(docx_path)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python md_to_word.py <path_to_markdown_file> [optional_output_path.docx]")
        sys.exit(1)
        
    input_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None
    convert_md_to_docx(input_path, out_path)
